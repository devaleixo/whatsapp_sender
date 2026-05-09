from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from ..deps import get_db
from ..models import Campaign, Contact
from ..schemas import ContactIn, ContactOut, ImportResult
from ..services.phone import to_e164


router = APIRouter(prefix="/campaigns/{campaign_id}/contacts", tags=["contacts"])


@router.get("", response_model=list[ContactOut])
def list_contacts(campaign_id: int, db: Session = Depends(get_db)):
    if not db.get(Campaign, campaign_id):
        raise HTTPException(404, "campaign not found")
    return db.query(Contact).filter(Contact.campaign_id == campaign_id).order_by(Contact.name).all()


@router.post("", response_model=ContactOut, status_code=201)
def create_contact(campaign_id: int, payload: ContactIn, db: Session = Depends(get_db)):
    if not db.get(Campaign, campaign_id):
        raise HTTPException(404, "campaign not found")
    e164 = to_e164(payload.phone)
    if not e164:
        raise HTTPException(400, "invalid phone")
    if db.query(Contact).filter_by(campaign_id=campaign_id, e164_phone=e164).first():
        raise HTTPException(409, "contact already exists in this campaign")
    c = Contact(campaign_id=campaign_id, e164_phone=e164, **payload.model_dump())
    db.add(c)
    db.commit()
    db.refresh(c)
    return c


@router.delete("/{contact_id}", status_code=204)
def delete_contact(campaign_id: int, contact_id: int, db: Session = Depends(get_db)):
    c = db.get(Contact, contact_id)
    if not c or c.campaign_id != campaign_id:
        raise HTTPException(404)
    db.delete(c)
    db.commit()


@router.post("/import", response_model=ImportResult)
async def import_xlsx(
    campaign_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    if not db.get(Campaign, campaign_id):
        raise HTTPException(404, "campaign not found")

    content = await file.read()
    try:
        wb = load_workbook(BytesIO(content), data_only=True)
    except Exception as e:
        raise HTTPException(400, f"invalid xlsx: {e}")

    ws = wb.active
    imported = skipped = invalid = 0

    existing = set(
        r[0] for r in db.execute(
            Contact.__table__.select()
            .with_only_columns(Contact.e164_phone)
            .where(Contact.campaign_id == campaign_id)
        ).all()
    )

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if header_row is None:
        return ImportResult(imported=0, skipped=0, invalid=0)

    # normalize header names: lowercase + strip accents for flexible matching
    import unicodedata

    def _norm(s: str) -> str:
        s = unicodedata.normalize("NFD", str(s).lower().strip())
        return "".join(c for c in s if unicodedata.category(c) != "Mn")

    col = {_norm(h): i for i, h in enumerate(header_row) if h is not None}

    def _get(row, *keys):
        for k in keys:
            idx = col.get(_norm(k))
            if idx is not None and idx < len(row):
                return row[idx]
        return None

    for row in ws.iter_rows(min_row=2, values_only=True):
        name = _get(row, "Nome", "Name")
        phone = _get(row, "Telefone", "Phone")
        address = _get(row, "Endereco", "Endereço", "Address")
        neighborhood = _get(row, "Bairro", "Neighborhood")
        rating = _get(row, "Avaliacao", "Avaliação", "Rating")
        rating_count = _get(row, "Qtd_Avaliacoes", "Qtd Avaliacoes", "Rating Count")
        website = _get(row, "Website")
        business_type = _get(row, "Tipo", "Type")
        business_status = _get(row, "Status")
        place_id_raw = _get(row, "Place_ID", "Place ID")

        if not name or not phone or str(phone) == "N/A":
            invalid += 1
            continue
        e164 = to_e164(str(phone))
        if not e164:
            invalid += 1
            continue
        if e164 in existing:
            skipped += 1
            continue

        place_id = str(place_id_raw) if place_id_raw else None

        db.add(Contact(
            campaign_id=campaign_id,
            name=str(name),
            phone=str(phone),
            e164_phone=e164,
            address=str(address) if address else None,
            neighborhood=str(neighborhood) if neighborhood else None,
            rating=str(rating) if rating else None,
            rating_count=str(rating_count) if rating_count else None,
            website=str(website) if website else None,
            business_type=str(business_type) if business_type else None,
            business_status=str(business_status) if business_status else None,
            place_id=place_id,
            source=place_id or "xlsx",
        ))
        existing.add(e164)
        imported += 1

    db.commit()
    return ImportResult(imported=imported, skipped=skipped, invalid=invalid)
